#!/usr/bin/env python3
"""
Sprint Report Generator
Generates a formatted Excel sprint report from Jira ticket data (JSON input).

Usage:
    python generate_sprint_report.py --data tickets.json --output "Sprint 3-1-26 ~ 3-14-26.xlsx"
    python generate_sprint_report.py --data tickets.json --output report.xlsx --team-config sprint_team_config.md
"""

import json
import re
import sys
import argparse
from datetime import datetime, date
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)

YELLOW_FILL = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
PINK_FILL = PatternFill(start_color="FFFFCCFF", end_color="FFFFCCFF", fill_type="solid")
RED_FILL = PatternFill(start_color="FFFF6600", end_color="FFFF6600", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")
GREEN_FILL = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")
COND_RED_FILL = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")

HEADER_FONT = Font(bold=True, size=11)

COLUMNS = [
    ("Finished Date",            15.57),
    ("Create Date",              15.71),
    ("Started Date",             14.29),
    ("On Hold Date",             15.00),
    ("Re-Started Date",          17.43),
    ("Assignation Date",         18.29),
    ("Label",                    17.00),
    ("Estimation ",              13.14),
    ("Completed",                11.00),
    ("Non Workable Days",        21.00),
    ("Real Worked time",         19.57),
    ("Difference est and real",  23.86),
    ("Assigned By",              15.43),
    ("Dev",                      15.14),
    ("Project",                  18.43),
    ("Jira Ticket Number",       20.00),
    ("Ticket Name",              80.00),
    ("Resolution Details",       62.00),
    ("Status",                   17.57),
    ("Effort Comments",          54.86),
]

STATUS_ORDER = {"done": 0, "to do": 1, "in progress": 2}

DATE_FORMATS = [
    "%Y-%m-%dT%H:%M:%S.%f",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d",
]


def parse_date(date_str):
    if not date_str:
        return None
    cleaned = re.sub(r"[+-]\d{2}:?\d{2}$", "", date_str).replace("Z", "")
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(cleaned, fmt)
        except ValueError:
            continue
    return None


def to_date(dt_or_str):
    """Convert a datetime, date, or date string to a date object."""
    if dt_or_str is None:
        return None
    if isinstance(dt_or_str, datetime):
        return dt_or_str.date()
    if isinstance(dt_or_str, date):
        return dt_or_str
    parsed = parse_date(str(dt_or_str))
    return parsed.date() if parsed else None


def format_short_date(dt):
    """Format datetime as M-D-YY (no leading zeros)."""
    return f"{dt.month}-{dt.day}-{str(dt.year)[-2:]}"


def get_row_fill(ticket):
    status = (ticket.get("status") or "").strip().lower()
    has_finished = bool(ticket.get("finished_date"))

    if has_finished or status == "done":
        return YELLOW_FILL
    if status == "to do":
        return PINK_FILL
    if status == "in progress":
        return RED_FILL
    return YELLOW_FILL


# ---------------------------------------------------------------------------
# Team-config markdown parser
# ---------------------------------------------------------------------------

def _parse_md_table(lines):
    """Yield dicts for each data row in a markdown table (header + separator + rows)."""
    table_rows = [ln for ln in lines if ln.strip().startswith("|")]
    if len(table_rows) < 2:
        return
    headers = [h.strip() for h in table_rows[0].strip().strip("|").split("|")]
    for line in table_rows[1:]:
        if re.match(r"^\s*\|[\s\-:|]+\|\s*$", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) >= len(headers):
            yield dict(zip(headers, cells))


def parse_team_config(md_path):
    """
    Parse sprint_team_config.md and return:
      {
        "members": [{"Name": ..., "Email": ...}, ...],
        "days_off": [{"Name": ..., "Date": ..., "Reason": ...}, ...],
        "holidays": [{"Date": ..., "Reason": ...}, ...],
        "cloud_id": "...",
        "project": "...",
      }
    """
    text = Path(md_path).read_text(encoding="utf-8")
    sections = re.split(r"^## ", text, flags=re.MULTILINE)

    result = {
        "members": [],
        "days_off": [],
        "holidays": [],
        "cloud_id": "eci-solutions.atlassian.net",
        "project": "MFGR10",
    }

    for section in sections:
        heading = section.split("\n", 1)[0].strip().lower()
        section_lines = section.split("\n")

        if heading.startswith("team members"):
            result["members"] = list(_parse_md_table(section_lines))
        elif heading.startswith("days off"):
            result["days_off"] = list(_parse_md_table(section_lines))
        elif heading.startswith("team holidays"):
            result["holidays"] = list(_parse_md_table(section_lines))
        elif heading.startswith("jira settings"):
            cloud_match = re.search(r"\*\*Cloud ID\*\*:\s*(.+)", section)
            proj_match = re.search(r"\*\*Project\*\*:\s*(.+)", section)
            if cloud_match:
                result["cloud_id"] = cloud_match.group(1).strip()
            if proj_match:
                result["project"] = proj_match.group(1).strip()

    return result


def compute_non_workable_days(ticket, team_cfg, sprint_end_date):
    """
    Count how many off-days/holidays overlap with a ticket's work period.
    Work period = started_date .. finished_date (or sprint_end if not finished).
    """
    start = to_date(ticket.get("started_date"))
    end = to_date(ticket.get("finished_date")) or sprint_end_date
    if not start or not end:
        return 0

    dev_name = (ticket.get("dev") or "").strip().lower()
    off_dates = set()

    for entry in team_cfg.get("holidays", []):
        d = to_date(entry.get("Date"))
        if d:
            off_dates.add(d)

    for entry in team_cfg.get("days_off", []):
        if (entry.get("Name") or "").strip().lower() == dev_name:
            d = to_date(entry.get("Date"))
            if d:
                off_dates.add(d)

    return sum(1 for d in off_dates if start <= d <= end)


# ---------------------------------------------------------------------------
# Report generation
# ---------------------------------------------------------------------------

def generate_report(data, output_path, team_cfg=None):
    wb = Workbook()

    sprint_start = data["sprint_start"]
    sprint_end = data["sprint_end"]
    start_dt = parse_date(sprint_start)
    end_dt = parse_date(sprint_end)
    sprint_end_date = to_date(sprint_end)

    sheet_label = (
        f"Sprint {format_short_date(start_dt)} | {format_short_date(end_dt)}"
        if start_dt and end_dt
        else "Sprint Report"
    )
    if len(sheet_label) > 31:
        sheet_label = sheet_label[:31]

    ws = wb.active
    ws.title = sheet_label

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, (header, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    sprint_start_date = to_date(sprint_start)

    all_tickets = data.get("tickets", [])
    filtered = [
        t for t in all_tickets
        if (to_date(t.get("started_date")) or sprint_start_date) >= sprint_start_date
    ]

    tickets = sorted(
        filtered,
        key=lambda t: (
            (t.get("dev") or "").lower(),
            STATUS_ORDER.get((t.get("status") or "").strip().lower(), 3),
            t.get("started_date") or "",
        ),
    )

    if team_cfg:
        for ticket in tickets:
            if ticket.get("non_workable_days", 0) == 0:
                ticket["non_workable_days"] = compute_non_workable_days(
                    ticket, team_cfg, sprint_end_date
                )

    for row_idx, ticket in enumerate(tickets, 2):
        fill = get_row_fill(ticket)
        r = row_idx

        ws.cell(row=r, column=1, value=parse_date(ticket.get("finished_date")))
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"

        ws.cell(row=r, column=2, value=parse_date(ticket.get("create_date")))
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd hh:mm"

        started_val = parse_date(ticket.get("started_date"))
        ws.cell(row=r, column=3, value=started_val)
        ws.cell(row=r, column=3).number_format = "yyyy-mm-dd"
        if started_val is None:
            ws.cell(row=r, column=3).fill = RED_FILL

        ws.cell(row=r, column=4, value=parse_date(ticket.get("on_hold_date")))
        ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"

        ws.cell(row=r, column=5, value=parse_date(ticket.get("restarted_date")))
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"

        ws.cell(row=r, column=6, value=parse_date(ticket.get("assignation_date")))
        ws.cell(row=r, column=6).number_format = "yyyy-mm-dd"

        ws.cell(row=r, column=7, value=ticket.get("label", ""))
        ws.cell(row=r, column=8, value=ticket.get("estimation_hours", 0))

        ws.cell(
            row=r, column=9,
            value=f"=IF(ISBLANK(A{r}),0,(A{r}-C{r})*8)",
        )

        ws.cell(row=r, column=10, value=ticket.get("non_workable_days", 0))

        on_hold_hours = (
            f"IF(ISBLANK(D{r}),0,"
            f"IF(NOT(ISBLANK(E{r})),(E{r}-D{r})*8,"
            f"IF(NOT(ISBLANK(A{r})),(A{r}-D{r})*8,0)))"
        )
        ws.cell(row=r, column=11, value=f"=IF(ISBLANK(C{r}),0,I{r}-{on_hold_hours}-(J{r}*8))")
        ws.cell(row=r, column=11).number_format = "0"
        ws.cell(row=r, column=12, value=f"=IF(K{r}>0,H{r}-K{r},0)")
        ws.cell(row=r, column=12).number_format = "0"

        ws.cell(row=r, column=13, value=ticket.get("assigned_by", ""))
        ws.cell(row=r, column=14, value=ticket.get("dev", ""))
        ws.cell(row=r, column=15, value=ticket.get("project", ""))
        ws.cell(row=r, column=16, value=ticket.get("jira_ticket", ""))
        ws.cell(row=r, column=17, value=ticket.get("ticket_name", ""))
        ws.cell(row=r, column=18, value=ticket.get("resolution_details", ""))
        ws.cell(row=r, column=19, value=ticket.get("status", ""))
        ws.cell(row=r, column=20, value=ticket.get("effort_comments", ""))

        for col in range(1, 21):
            cell = ws.cell(row=r, column=col)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=(col >= 17))
            cell.border = thin_border

        ws.cell(row=r, column=12).fill = WHITE_FILL

    ws.freeze_panes = "A2"
    last_row = len(tickets) + 1
    ws.auto_filter.ref = f"A1:T{last_row}"

    if last_row > 1:
        diff_range = f"L2:L{last_row}"
        ws.conditional_formatting.add(
            diff_range,
            CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN_FILL),
        )
        ws.conditional_formatting.add(
            diff_range,
            CellIsRule(operator="lessThan", formula=["0"], fill=COND_RED_FILL),
        )

    wb.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Generate Sprint Report Excel")
    parser.add_argument("--data", required=True, help="Path to JSON data file with ticket info")
    parser.add_argument("--output", required=True, help="Output .xlsx file path")
    parser.add_argument(
        "--team-config",
        default=None,
        help="Path to sprint_team_config.md for auto-calculating non-workable days",
    )
    args = parser.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    team_cfg = None
    if args.team_config:
        team_cfg = parse_team_config(args.team_config)

    result = generate_report(data, args.output, team_cfg)
    print(f"Report generated: {result}")


if __name__ == "__main__":
    main()
